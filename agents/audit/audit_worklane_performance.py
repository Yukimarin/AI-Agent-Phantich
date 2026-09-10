import json
import os
import re
import unicodedata
import openpyxl
import sys

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

def normalize_vietnamese_name(name):
    if not name:
        return ""
    name = " ".join(name.strip().split())
    name = name.lower()
    name = unicodedata.normalize('NFKD', name)
    name = "".join([c for c in name if not unicodedata.combining(c)])
    name = name.replace("đ", "d")
    return name

SPECIAL_MAPPINGS = {
    "lưu xuân hoàng nguyên": "lưu hoàng xuân nguyên",
    "xuân nguyên": "lưu hoàng xuân nguyên"
}

def normalize_name(name):
    norm = name.strip().lower()
    if norm in SPECIAL_MAPPINGS:
        norm = SPECIAL_MAPPINGS[norm]
    return norm

LEAVE_DAYS_AUG = {
    "nguyễn thị như quỳnh": ["2026-08-10"],
    "nguyễn thị tươi": ["2026-08-13"],
    "nguyễn ngọc vân khanh": ["2026-08-12"],
    "trần minh cường": ["2026-08-14"]
}

WORKDAYS_AUG = [
    "2026-08-03", "2026-08-04", "2026-08-05", "2026-08-06", "2026-08-07",
    "2026-08-10", "2026-08-11", "2026-08-12", "2026-08-13", "2026-08-14",
    "2026-08-17", "2026-08-18", "2026-08-19", "2026-08-20", "2026-08-21",
    "2026-08-24", "2026-08-25", "2026-08-26", "2026-08-27", "2026-08-28"
]

WORKDAYS_SEPT = [
    "2026-09-03", "2026-09-04", "2026-09-07", "2026-09-08"
]

ALL_DATES_SEPT = [
    "2026-09-01", "2026-09-02", "2026-09-03", "2026-09-04",
    "2026-09-05", "2026-09-06", "2026-09-07", "2026-09-08"
]

EXACT_STAFF_ROSTER = {
    # 1. Khối CNTT Hà Nội - Leader: Hồ Xuân Hùng (Rank 5)
    "hồ xuân hùng": {"name": "Hồ Xuân Hùng", "role": "Leader", "rank": 5, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hồ Xuân Hùng"},
    "lương quốc tuấn": {"name": "Lương Quốc Tuấn", "role": "Giảng viên", "rank": 3, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hồ Xuân Hùng"},
    "nguyễn quảng an": {"name": "Nguyễn Quảng An", "role": "Giảng viên", "rank": 4, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hồ Xuân Hùng"},
    "lâm tùng dương": {"name": "Lâm Tùng Dương", "role": "Giảng viên", "rank": 3, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hồ Xuân Hùng"},
    "ngọ văn quý": {"name": "Ngọ Văn Quý", "role": "Giảng viên", "rank": 4, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hồ Xuân Hùng"},
    "lại trung lâm": {"name": "Lại Trung Lâm", "role": "Trợ giảng", "rank": 2, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hồ Xuân Hùng"},
    "phạm ngọc kiên": {"name": "Phạm Ngọc Kiên", "role": "Trợ giảng", "rank": 2, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hồ Xuân Hùng"},

    "bùi thanh hải": {"name": "Bùi Thanh Hải", "role": "Giảng viên", "rank": 5, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - HPC", "leader": "Hồ Xuân Hùng"},
    "trịnh quốc hai": {"name": "Trịnh Quốc Hai", "role": "Giảng viên", "rank": 4, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - HPC", "leader": "Hồ Xuân Hùng"},
    "nguyễn công hưởng": {"name": "Nguyễn Công Hưởng", "role": "Giảng viên", "rank": 3, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - HPC", "leader": "Hồ Xuân Hùng"},
    "phạm tuấn bình": {"name": "Phạm Tuấn Bình", "role": "Giảng viên", "rank": 3, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - HPC", "leader": "Hồ Xuân Hùng"},
    "đinh thành nam": {"name": "Đinh Thành Nam", "role": "Trợ giảng", "rank": 2, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - HPC", "leader": "Hồ Xuân Hùng"},
    "mai xuân chinh": {"name": "Mai Xuân Chinh", "role": "Trợ giảng", "rank": 2, "group": "Khối CNTT Hà Nội", "campus": "Cơ sở Hà Nội - HPC", "leader": "Hồ Xuân Hùng"},

    # 2. Khối QTKD - Leader: Hoàng Thị Kim Oanh (Rank 5)
    "hoàng thị kim oanh": {"name": "Hoàng Thị Kim Oanh", "role": "Leader", "rank": 5, "group": "Khối Quản trị Kinh doanh", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hoàng Thị Kim Oanh"},
    "hoàng thị hậu": {"name": "Hoàng Thị Hậu", "role": "Giảng viên", "rank": 5, "group": "Khối Quản trị Kinh doanh", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hoàng Thị Kim Oanh"},
    "đặng quỳnh trang": {"name": "Đặng Quỳnh Trang", "role": "Giảng viên", "rank": 3, "group": "Khối Quản trị Kinh doanh", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hoàng Thị Kim Oanh"},
    "nguyễn ngọc vân khanh": {"name": "Nguyễn Ngọc Vân Khanh", "role": "Giảng viên", "rank": 3, "group": "Khối Quản trị Kinh doanh", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hoàng Thị Kim Oanh"},
    "nguyễn thị hồng minh": {"name": "Nguyễn Thị Hồng Minh", "role": "Giảng viên", "rank": 3, "group": "Khối Quản trị Kinh doanh", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hoàng Thị Kim Oanh"},
    "nguyễn thị như quỳnh": {"name": "Nguyễn Thị Như Quỳnh", "role": "Trợ giảng", "rank": 1, "group": "Khối Quản trị Kinh doanh", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hoàng Thị Kim Oanh"},
    "triệu thị thanh tâm": {"name": "Triệu Thị Thanh Tâm", "role": "Trợ giảng", "rank": 1, "group": "Khối Quản trị Kinh doanh", "campus": "Cơ sở Hà Nội - Ngọc Trục", "leader": "Hoàng Thị Kim Oanh"},
    "lê nhựt mi": {"name": "Lê Nhựt Mi", "role": "Giảng viên", "rank": 3, "group": "Khối Quản trị Kinh doanh", "campus": "Cơ sở Hồ Chí Minh", "leader": "Hoàng Thị Kim Oanh"},
    "lê thị bảo yến": {"name": "Lê Thị Bảo Yến", "role": "Trợ giảng", "rank": 2, "group": "Khối Quản trị Kinh doanh", "campus": "Cơ sở Hồ Chí Minh", "leader": "Hoàng Thị Kim Oanh"},

    # 3. Khối Ngoại ngữ & KNM - Leader: Giáp Thị Minh Hằng (Rank 5)
    "giáp thị minh hằng": {"name": "Giáp Thị Minh Hằng", "role": "Leader", "rank": 5, "group": "Khối Ngoại ngữ và KNM", "campus": "Cơ sở Hà Nội - HPC", "leader": "Giáp Thị Minh Hằng"},
    "lò thị ngọc anh": {"name": "Lò Thị Ngọc Anh", "role": "Giảng viên", "rank": 5, "group": "Khối Ngoại ngữ và KNM", "campus": "Cơ sở Hà Nội - HPC", "leader": "Giáp Thị Minh Hằng"},
    "lê thị đỏ": {"name": "Lê Thị Đỏ", "role": "Giảng viên", "rank": 3, "group": "Khối Ngoại ngữ và KNM", "campus": "Cơ sở Hà Nội - HPC", "leader": "Giáp Thị Minh Hằng"},
    "ngô quang huấn": {"name": "Ngô Quang Huấn", "role": "Giảng viên", "rank": 3, "group": "Khối Ngoại ngữ và KNM", "campus": "Cơ sở Hà Nội - HPC", "leader": "Giáp Thị Minh Hằng"},

    # 4. Khối QLCLĐT Hà Nội - Leader: Nguyễn Thị Tươi (Rank 4)
    "nguyễn thị tươi": {"name": "Nguyễn Thị Tươi", "role": "Leader", "rank": 4, "group": "Khối QLCLĐT Hà Nội", "campus": "Cơ sở Hà Nội - HPC", "leader": "Nguyễn Thị Tươi"},
    "nguyễn huyền trang": {"name": "Nguyễn Huyền Trang", "role": "Giáo vụ", "rank": 2, "group": "Khối QLCLĐT Hà Nội", "campus": "Cơ sở Hà Nội - HPC", "leader": "Nguyễn Thị Tươi"},
    "trần thị mỹ phước": {"name": "Trần Thị Mỹ Phước", "role": "Giáo vụ", "rank": 2, "group": "Khối QLCLĐT Hà Nội", "campus": "Cơ sở Hồ Chí Minh", "leader": "Nguyễn Thị Tươi"},
    "nguyễn xuân bách": {"name": "Nguyễn Xuân Bách", "role": "Giảng viên", "rank": 4, "group": "Khối QLCLĐT Hà Nội", "campus": "Cơ sở Hà Nội - HPC", "leader": "Nguyễn Thị Tươi"},

    # 5. Khối CNTT HCM - Leader: Nguyễn Bá Minh Đạo (Rank 5)
    "nguyễn bá minh đạo": {"name": "Nguyễn Bá Minh Đạo", "role": "Leader", "rank": 5, "group": "Khối CNTT HCM", "campus": "Cơ sở Hồ Chí Minh", "leader": "Nguyễn Bá Minh Đạo"},
    "lê hà thanh sang": {"name": "Lê Hà Thanh Sang", "role": "Giảng viên", "rank": 4, "group": "Khối CNTT HCM", "campus": "Cơ sở Hồ Chí Minh", "leader": "Nguyễn Bá Minh Đạo"},
    "lưu hoàng xuân nguyên": {"name": "Lưu Hoàng Xuân Nguyên", "role": "Trợ giảng", "rank": 2, "group": "Khối CNTT HCM", "campus": "Cơ sở Hồ Chí Minh", "leader": "Nguyễn Bá Minh Đạo"},
    "phạm viết hùng": {"name": "Phạm Viết Hùng", "role": "Trợ giảng", "rank": 2, "group": "Khối CNTT HCM", "campus": "Cơ sở Hồ Chí Minh", "leader": "Nguyễn Bá Minh Đạo"},
    "trần quốc tuấn": {"name": "Trần Quốc Tuấn", "role": "Giảng viên", "rank": 3, "group": "Khối CNTT HCM", "campus": "Cơ sở Hồ Chí Minh", "leader": "Nguyễn Bá Minh Đạo"},
    "đặng minh luân": {"name": "Đặng Minh Luân", "role": "Trợ giảng thử việc", "rank": 1, "group": "Khối CNTT HCM", "campus": "Cơ sở Hồ Chí Minh", "leader": "Nguyễn Bá Minh Đạo"},
    "nguyễn đức minh": {"name": "Nguyễn Đức Minh", "role": "Giảng viên", "rank": 4, "group": "Khối CNTT HCM", "campus": "Cơ sở Hồ Chí Minh", "leader": "Nguyễn Bá Minh Đạo"},
    "nguyễn ngọc sơn": {"name": "Nguyễn Ngọc Sơn", "role": "Trợ giảng thử việc", "rank": 1, "group": "Khối CNTT HCM", "campus": "Cơ sở Hồ Chí Minh", "leader": "Nguyễn Bá Minh Đạo"},
    "phan ngọc tài": {"name": "Phan Ngọc Tài", "role": "Trợ giảng thử việc", "rank": 1, "group": "Khối CNTT HCM", "campus": "Cơ sở Hồ Chí Minh", "leader": "Nguyễn Bá Minh Đạo"},

    # 6. LMS AI - Leader: Trần Minh Cường (Rank 5)
    "trần minh cường": {"name": "Trần Minh Cường", "role": "Leader", "rank": 5, "group": "LMS AI", "campus": "Cơ sở Hà Nội - HPC", "leader": "Trần Minh Cường"},
    "lê thành ngọc": {"name": "Lê Thành Ngọc", "role": "Giảng viên", "rank": 3, "group": "LMS AI", "campus": "Cơ sở Hồ Chí Minh", "leader": "Trần Minh Cường"}
}

def load_staff_database():
    return EXACT_STAFF_ROSTER.copy()

def load_foreign_lang_kpi_master():
    nn_standards = []
    nn_mappings = {}
    nn_path = r"C:\Users\DELL\Downloads\[ENG&JPN] KPI MASTER KHỐI ĐÀO TẠO NGOẠI NGỮ.xlsx"
    if not os.path.exists(nn_path):
        nn_path = r"output/reports/KPI_Master_Khoi_Ngoai_Ngu.xlsx"
        
    if os.path.exists(nn_path):
        try:
            wb = openpyxl.load_workbook(nn_path, data_only=True)
            if "IMPORT_WORKLANE_DROPDOWN" in wb.sheetnames:
                ws = wb["IMPORT_WORKLANE_DROPDOWN"]
                for r in range(2, ws.max_row + 1):
                    title = ws.cell(r, 2).value
                    role = ws.cell(r, 3).value
                    rank = ws.cell(r, 4).value
                    mins = ws.cell(r, 5).value
                    unit = ws.cell(r, 7).value
                    cat = ws.cell(r, 8).value
                    key = ws.cell(r, 9).value
                    if title and mins is not None:
                        nn_standards.append({
                            "name": str(title).strip().lower(),
                            "role": str(role).strip().lower() if role else "",
                            "rank": int(rank) if rank is not None and str(rank).isdigit() else 3,
                            "minutes": float(mins),
                            "unit": str(unit).strip().lower() if unit else "",
                            "category": str(cat).strip() if cat else "",
                            "key": str(key).strip().lower() if key else ""
                        })
            if "MAPPING_WORKLANE_TASKS" in wb.sheetnames:
                ws_m = wb["MAPPING_WORKLANE_TASKS"]
                for r in range(2, ws_m.max_row + 1):
                    code = ws_m.cell(r, 7).value
                    task_title = ws_m.cell(r, 8).value
                    master_key = ws_m.cell(r, 12).value
                    std_name = ws_m.cell(r, 13).value
                    if task_title:
                        nn_mappings[str(task_title).strip().lower()] = {
                            "code": str(code).strip() if code else "",
                            "key": str(master_key).strip().lower() if master_key else "",
                            "std_name": str(std_name).strip() if std_name else ""
                        }
            wb.close()
            print(f"Đã nạp thành công {len(nn_standards)} định mức Ngoại ngữ & {len(nn_mappings)} mappings từ: {os.path.basename(nn_path)}")
        except Exception as e:
            print("Lỗi đọc Ngoại ngữ KPI Master:", e)
    return nn_standards, nn_mappings

def load_kpi_masters_v2():
    qtkd_master = []
    cntt_master = []
    
    # 1. QTKD KPI Master (Ưu tiên bản Draft 8/2026 mới nhất)
    qtkd_new_path = r"C:\Users\DELL\Downloads\[QTKDS] KPI Master (Draft - 8_2026).xlsx"
    qtkd_old_path = r"C:\Users\DELL\Downloads\_Task Management_ QL Khối QTKD (3).xlsx"
    qtkd_path = qtkd_new_path if os.path.exists(qtkd_new_path) else qtkd_old_path
    
    if os.path.exists(qtkd_path):
        try:
            wb = openpyxl.load_workbook(qtkd_path, data_only=True)
            if "Sheet1" in wb.sheetnames:
                ws = wb["Sheet1"]
                for r in range(2, ws.max_row + 1):
                    name = ws.cell(r, 1).value
                    role = ws.cell(r, 2).value
                    rank = ws.cell(r, 3).value
                    unit = ws.cell(r, 4).value
                    mins_old = ws.cell(r, 5).value
                    mins_new = ws.cell(r, 6).value
                    mins = mins_new if mins_new is not None else mins_old
                    if name and mins is not None:
                        qtkd_master.append({
                            "name": str(name).strip().lower(),
                            "role": str(role).strip().lower() if role else "",
                            "rank": int(rank) if rank is not None and str(rank).isdigit() else 3,
                            "unit": str(unit).strip().lower() if unit else "",
                            "minutes": float(mins)
                        })
            elif "KPI_MASTER" in wb.sheetnames:
                ws = wb["KPI_MASTER"]
                for r in range(2, ws.max_row + 1):
                    name = ws.cell(r, 1).value
                    ttype = ws.cell(r, 2).value
                    role = ws.cell(r, 3).value
                    rank = ws.cell(r, 4).value
                    unit = ws.cell(r, 5).value
                    mins = ws.cell(r, 6).value
                    if name and mins is not None:
                        qtkd_master.append({
                            "name": str(name).strip().lower(),
                            "task_type": str(ttype).strip().lower() if ttype else "",
                            "role": str(role).strip().lower() if role else "",
                            "rank": int(rank) if rank is not None and str(rank).isdigit() else 3,
                            "unit": str(unit).strip().lower() if unit else "",
                            "minutes": float(mins)
                        })
            wb.close()
            print(f"Đã nạp thành công {len(qtkd_master)} định mức QTKD từ: {os.path.basename(qtkd_path)}")
        except Exception as e:
            print("Lỗi đọc QTKD KPI Master:", e)
            
    # 2. CNTT KPI Master (Ưu tiên bản FINAL mới nhất)
    cntt_final_path = r"C:\Users\DELL\Downloads\KPI_MASTER_Giang_vien_Tro_giang_FINAL.xlsx"
    cntt_path_old = r"C:\Users\DELL\Downloads\Quản lý hiệu suất đào tạo (2).xlsx"
    cntt_path_alt = r"C:\Users\DELL\Downloads\Quản lý hiệu suất đào tạo.xlsx"
    
    if os.path.exists(cntt_final_path):
        try:
            wb = openpyxl.load_workbook(cntt_final_path, data_only=True)
            if "KPI_MASTER" in wb.sheetnames:
                ws = wb["KPI_MASTER"]
                for r in range(2, ws.max_row + 1):
                    kpi_name = ws.cell(r, 1).value
                    ttype = ws.cell(r, 2).value
                    role = ws.cell(r, 3).value
                    rank = ws.cell(r, 4).value
                    mins = ws.cell(r, 5).value
                    key = ws.cell(r, 6).value
                    if kpi_name and mins is not None:
                        cntt_master.append({
                            "name": str(kpi_name).strip().lower(),
                            "task_type": str(ttype).strip().lower() if ttype else "",
                            "role": str(role).strip().lower() if role else "",
                            "rank": int(rank) if rank is not None and str(rank).isdigit() else 3,
                            "minutes": float(mins),
                            "key": str(key).strip().lower() if key else ""
                        })
            wb.close()
            print(f"Đã nạp thành công {len(cntt_master)} định mức CNTT từ: {os.path.basename(cntt_final_path)}")
        except Exception as e:
            print("Lỗi đọc CNTT KPI Master FINAL:", e)
    elif os.path.exists(cntt_path_old) or os.path.exists(cntt_path_alt):
        target_path = cntt_path_old if os.path.exists(cntt_path_old) else cntt_path_alt
        try:
            wb = openpyxl.load_workbook(target_path, data_only=True)
            sheetname = "Cấu trúc KPI công việc GV. TG"
            if sheetname in wb.sheetnames:
                ws = wb[sheetname]
                for r in range(2, ws.max_row + 1):
                    name = ws.cell(r, 2).value
                    role = ws.cell(r, 3).value
                    rank = ws.cell(r, 4).value
                    mins = ws.cell(r, 5).value
                    cat = ws.cell(r, 6).value
                    if name and mins is not None:
                        cntt_master.append({
                            "name": str(name).strip().lower(),
                            "role": str(role).strip().lower() if role else "",
                            "rank": int(rank) if rank is not None and str(rank).isdigit() else 3,
                            "minutes": float(mins),
                            "category": str(cat).strip().lower() if cat else ""
                        })
            wb.close()
            print(f"Đã nạp fallback {len(cntt_master)} định mức CNTT từ: {os.path.basename(target_path)}")
        except Exception as e:
            print("Lỗi đọc CNTT KPI Master fallback:", e)
            
    return qtkd_master, cntt_master

def match_foreign_lang_standard(role, rank, title, nn_standards, nn_mappings):
    t_low = title.lower().strip()
    clean_t = re.sub(r'\s*\(\d{2}/\d{2}\)', '', t_low).strip()
    clean_t = re.sub(r'^\[re\]\s*', '', clean_t).strip()
    clean_t = re.sub(r'^\[rexjax\]\s*', '', clean_t).strip()

    # 1. Tra cứu trực tiếp từ bảng ánh xạ 284 tasks Worklane Ngoại ngữ
    for map_t, m_val in nn_mappings.items():
        map_clean = re.sub(r'\s*\(\d{2}/\d{2}\)', '', map_t).strip()
        map_clean = re.sub(r'^\[re\]\s*', '', map_clean).strip()
        map_clean = re.sub(r'^\[rexjax\]\s*', '', map_clean).strip()
        if map_clean in clean_t or clean_t in map_clean or map_t in t_low or t_low in map_t:
            target_key_prefix = m_val['key']
            for s in nn_standards:
                if s['key'].startswith(target_key_prefix) and s['rank'] == rank:
                    return s['minutes'] / 60.0, f"KPI Master NN: {s['name'].title()} ({s['key'].upper()})", False
            for s in nn_standards:
                if s['key'].startswith(target_key_prefix) and abs(s['rank'] - rank) <= 1:
                    return s['minutes'] / 60.0, f"KPI Master NN: {s['name'].title()} ({s['key'].upper()})", False
            return 1.0, f"Mapped NN: {m_val['std_name']}", False

    # 2. Quy tắc nghiệp vụ chuyên sâu theo 118 định mức Barem Ngoại ngữ
    if any(k in t_low for k in ["báo cáo ngày", "kế hoạch", "lập kế hoạch"]):
        return 0.5, "Báo cáo ngày & Lập kế hoạch", False

    if any(k in t_low for k in ["giao ban", "họp bộ môn", "họp viện", "họp triển khai", "họp giao ban"]):
        std = next((s for s in nn_standards if "nn-qt-giaoban" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else (1.5 if rank >= 5 else 2.0)), "Họp giao ban bộ môn / Viện", False

    if any(k in t_low for k in ["tuyển dụng", "phỏng vấn", "duyệt cv", "trao đổi thông tin và duyệt cv"]):
        std = next((s for s in nn_standards if "nn-qt-tuyendung" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else 1.0), "Rà soát CV & Phỏng vấn tuyển dụng", False

    if any(k in t_low for k in ["demo", "dự giờ", "teaching demo", "feedback gv"]):
        std = next((s for s in nn_standards if "nn-qt-demofb" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else (1.2 if rank >= 5 else 2.0)), "Dự giờ & Feedback GV", False

    if any(k in t_low for k in ["khung ctđt", "khung chương trình", "syllabus", "xây dựng ctđt"]):
        std = next((s for s in nn_standards if "nn-rnd-ctdt" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else (12.0 if rank >= 5 else 20.0)), "Xây dựng khung CTĐT môn học", False

    if any(k in t_low for k in ["timeline"]):
        std = next((s for s in nn_standards if "nn-rnd-timeline" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else (2.5 if rank >= 5 else 4.0)), "Xây dựng Timeline chi tiết", False

    if any(k in t_low for k in ["lesson plan", "giáo án"]):
        std = next((s for s in nn_standards if "nn-rnd-lessonplan" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else (1.0 if rank >= 5 else 2.0)), "Soạn Lesson Plan", False

    if any(k in t_low for k in ["slide", "bài giảng"]):
        std = next((s for s in nn_standards if "nn-hl-slide" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else (1.5 if rank >= 5 else 2.5)), "Soạn Slide bài giảng chuẩn", False

    if any(k in t_low for k in ["quy trình", "bàn giao ctđt", "bàn giao điều kiện", "review chương"]):
        std = next((s for s in nn_standards if "nn-rnd-quytrinh" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else (2.0 if rank >= 5 else 3.0)), "Chuẩn hóa quy trình đào tạo & Bàn giao CTĐT", False

    if any(k in t_low for k in ["lms ai", "order lms", "order tính năng", "test quy trình thiết lập"]):
        std = next((s for s in nn_standards if "nn-vh-orderai" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else (1.0 if rank >= 5 else 1.5)), "Order tính năng & Nghiệm thu LMS AI", False

    if any(k in t_low for k in ["duyệt đề", "check đề", "thẩm định đề", "rà soát đề"]):
        std = next((s for s in nn_standards if "nn-kt-duyetde" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else (1.0 if rank >= 5 else 1.5)), "Thẩm định, rà soát và Phê duyệt Đề thi", False

    if any(k in t_low for k in ["đề thi", "bộ đề", "bàn giao bộ đề", "tạo bộ câu hỏi", "quiz"]):
        std = next((s for s in nn_standards if "nn-kt-dethi" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else (2.0 if rank >= 5 else 3.0)), "Thiết kế Đề thi / Bộ đề kiểm tra", False

    if any(k in t_low for k in ["giảng dạy", "lên lớp", "dạy lớp talent", "triển khai lớp"]):
        std = next((s for s in nn_standards if "nn-gd-tructiep" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else 2.0), "Giảng dạy trên lớp chính khóa (120p)", False

    if any(k in t_low for k in ["onboard", "hướng dẫn tg", "hướng dẫn nhân sự"]):
        std = next((s for s in nn_standards if "nn-qt-onboard" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else 1.0), "Onboarding & Đào tạo hội nhập GV/TG", False

    if any(k in t_low for k in ["báo cáo", "tình hình thi", "tình hình nhân sự", "xử lý phản hồi", "xử lý học viên"]):
        std = next((s for s in nn_standards if "nn-qt-baocao" in s['key'] and s['rank'] == rank), None)
        return (std['minutes']/60.0 if std else 1.0), "Báo cáo định kỳ tiến độ & Quản trị", False

    if any(k in t_low for k in ["nghiên cứu", "học tập", "tự học", "bồi dưỡng"]):
        return 1.0, "Nghiên cứu / Học tập bồi dưỡng (Cần nghiệm thu)", True

    return 0.5, "Đầu việc tự do/chưa định mức", True

def match_kpi_standard(group, role, rank, title, qtkd_items, cntt_items, nn_standards, nn_mappings):
    """
    Thuật toán phân tích ngữ nghĩa đầu việc (Semantic Task Decomposition)
    Khớp chính xác với barem KPI Master theo Rank và Khối của nhân sự:
    """
    t_norm = title.lower().strip()
    is_qtkd = "qtkd" in group.lower() or "kinh doanh" in group.lower()
    is_nn = "ngoại ngữ" in group.lower() or "ngoaingu" in group.lower() or "tiếng" in group.lower()
    
    # 0. Nếu thuộc Khối Ngoại ngữ: Ưu tiên bộ 118 Barem và 284 Mapping của Khối Ngoại ngữ
    if is_nn:
        return match_foreign_lang_standard(role, rank, title, nn_standards, nn_mappings)

    items = qtkd_items if is_qtkd else cntt_items
    
    # 1. Thử tìm kiếm trực tiếp trong bảng KPI Master tương ứng của khối
    for item in items:
        name_k = item.get("name", "")
        if name_k and (name_k == t_norm or name_k in t_norm):
            item_rank = item.get("rank", 3)
            if item_rank == rank or abs(item_rank - rank) <= 1:
                return item["minutes"] / 60.0, f"KPI Master: {item.get('name', 'Chuẩn hóa')}", False

    # 2. Composite: Giảng dạy + Chuẩn bị giảng dạy
    if any(k in t_norm for k in ["giảng dạy", "lên lớp", "dạy lý thuyết"]) and any(k in t_norm for k in ["chuẩn bị", "soạn"]):
        prep_time = 0.5 if rank >= 4 else (0.75 if is_qtkd else 0.8)
        teach_time = 3.0 if is_qtkd else 2.0
        return (teach_time + prep_time), "Giảng dạy & Chuẩn bị lên lớp", False

    # 3. Giảng dạy độc lập
    if any(k in t_norm for k in ["giảng dạy", "lên lớp", "dạy trên lớp", "dạy lý thuyết", "triển khai buổi học"]):
        if is_qtkd:
            return 3.0, "Giảng dạy lý thuyết (180p)", False
        else:
            return 2.0, "Giảng dạy lý thuyết/thực hành (120p)", False

    # 4. Trợ giảng lên lớp / Thực hành
    if any(k in t_norm for k in ["hỗ trợ lớp", "trợ giảng lên lớp", "hỗ trợ thực hành", "hướng dẫn thực hành", "phụ đạo"]):
        if is_qtkd:
            return 3.0, "Triển khai buổi thực hành / TG lên lớp (180p)", False
        else:
            return 2.0, "Trợ giảng hỗ trợ thực hành (120p)", False

    # 5. Chuẩn bị giảng dạy / Giáo án / Kế hoạch bài giảng
    if any(k in t_norm for k in ["chuẩn bị giảng dạy", "soạn giáo án", "lesson plan", "soạn bài", "chuẩn bị bài"]):
        if is_qtkd:
            return 1.5, "Chuẩn bị giảng dạy & Giáo án (QTKD)", False
        else:
            return 0.5 if rank >= 4 else 0.8, "Chuẩn bị giảng dạy & Kế hoạch bài giảng", False

    # 6. Họp / Giao ban / Triển khai
    if any(k in t_norm for k in ["họp", "meeting", "giao ban", "sync", "trao đổi"]):
        if any(k in t_norm for k in ["giao ban", "chuyên môn", "định kỳ"]):
            return 2.0, "Họp chuyên môn / Giao ban định kỳ (120p)", False
        return 1.0, "Họp triển khai công việc (60p)", False

    # 7. Khảo thí / Chấm thi / Coi thi / Trông thi
    if any(k in t_norm for k in ["chấm thi", "chấm bài", "chấm assignment", "chấm lab", "chấm checkpoint", "chấm btvn", "trông thi", "coi thi"]):
        if any(k in t_norm for k in ["trông thi", "coi thi"]):
            return 1.0, "Trông thi / Coi thi sát hạch", False
        if is_qtkd:
            return 0.67, "Trợ giảng kiểm tra BTVN / Chấm bài (40p)", False
        else:
            return 2.0, "Chấm bài KT 30p đầu giờ / Chấm bài học phần (Khoán 120p/lớp)", False

    # 8. Trợ giảng CVHT / Chăm sóc học viên
    if any(k in t_norm for k in ["cvht", "chăm sóc học viên", "nhắc nhở học viên", "điểm danh", "liên hệ phụ huynh"]):
        if rank == 1:
            return 1.5, "Trợ giảng CVHT - Rank 1 (90p/ngày)", False
        else:
            return 2.0, "Trợ giảng CVHT - Rank 2 (120p/ngày)", False

    # 9. Triển khai 18 Task Review trong KPI Master FINAL CNTT
    if not is_qtkd and any(k in t_norm for k in ["review", "kiểm tra", "thẩm định", "nghiệm thu", "đánh giá", "duyệt"]):
        std_time = 0.5 if rank <= 2 else (0.4 if rank <= 4 else 0.3)
        return std_time, "Review học liệu độc lập (Master FINAL)", False

    # 10. Sản xuất học liệu trong CNTT Master FINAL
    if not is_qtkd and any(k in t_norm for k in ["làm slide", "soạn slide", "slide deck", "mindmap", "làm quiz", "xây dựng quiz", "tạo quiz", "viết bài đọc", "soạn bài đọc", "quay video"]):
        return 0.5, "Sản xuất học liệu ngoài KPI Master FINAL (Cần nghiệm thu)", True

    # 11. Báo cáo ngày / Kế hoạch công việc / Check thông báo
    if any(k in t_norm for k in ["báo cáo ngày", "kế hoạch", "lên kế hoạch", "báo cáo kết quả", "check mail", "check thông báo"]):
        return 0.5, "Báo cáo ngày & Lập kế hoạch", False

    # 12. Nghiên cứu / R&D / Học tập công nghệ / Bồi dưỡng
    if any(k in t_norm for k in ["nghiên cứu", "rnd", "tìm hiểu", "nccn", "test hệ thống", "học tập", "tập huấn", "bồi dưỡng", "ai agent", "ai day"]):
        return 1.0, "Nghiên cứu công nghệ / Học tập bồi dưỡng (Cần nghiệm thu)", True

    # 13. Thử tìm kiếm mờ mở rộng trong bảng KPI Master
    for item in items:
        name_k = item.get("name", "").lower()
        if name_k and (name_k in t_norm or t_norm in name_k):
            item_rank = item.get("rank", 3)
            if item_rank == rank or abs(item_rank - rank) <= 1:
                return item["minutes"] / 60.0, item.get("name", "KPI Master Match"), False

    # Task tự do chưa định mức
    return 0.5, "Đầu việc tự do/chưa định mức", True

def audit_period_data(period_name="sept_01_08"):
    staff_db = load_staff_database()
    qtkd_master, cntt_master = load_kpi_masters_v2()
    nn_standards, nn_mappings = load_foreign_lang_kpi_master()
    
    if period_name == "sept_01_08":
        cache_path = r"data/processed/daily_reports_sept_01_08.json"
        effective_workdays = WORKDAYS_SEPT
        all_dates_to_scan = ALL_DATES_SEPT
        period_title = "Kỳ 01/09 - 08/09/2026 (4 ngày làm việc)"
        audit_date_str = "2026-09-09"
        excess_hours_threshold = 6.0
        free_tasks_threshold = 2
    else:
        cache_path = r"data/processed/daily_reports_raw_cache.json"
        effective_workdays = WORKDAYS_AUG
        all_dates_to_scan = WORKDAYS_AUG
        period_title = "Tháng 08/2026 (20 ngày làm việc)"
        audit_date_str = "2026-09-04"
        excess_hours_threshold = 25.0
        free_tasks_threshold = 4

    if not os.path.exists(cache_path):
        print(f"Lỗi: Không tìm thấy file cache {cache_path}")
        return None

    with open(cache_path, "r", encoding="utf-8") as f:
        daily_cache = json.load(f)

    wl_issues_path = r"data/processed/project_issues_worklane.json"
    wl_issues = []
    if os.path.exists(wl_issues_path):
        with open(wl_issues_path, "r", encoding="utf-8") as f:
            wl_data = json.load(f)
            if isinstance(wl_data, dict):
                for p in wl_data.values():
                    iss_list = p.get('issues', {}).get('issues', [])
                    wl_issues.extend(iss_list)

    results = {}

    for norm_name, s_info in staff_db.items():
        name = s_info["name"]
        group = s_info["group"]
        role = s_info["role"]
        rank = s_info["rank"]
        campus = s_info["campus"]
        
        personal_leaves = LEAVE_DAYS_AUG.get(norm_name, []) if period_name == "august" else []
        staff_effective_days = [d for d in effective_workdays if d not in personal_leaves]
        expected_days_count = len(staff_effective_days)
        
        reported_days_count = 0
        total_declared_hours = 0.0
        total_standard_hours = 0.0
        
        excess_tasks = []
        free_floating_tasks = []
        unverified_tasks = []
        all_reported_tasks = []
        
        norm_vn = normalize_vietnamese_name(name)
        user_issues = [iss for iss in wl_issues if norm_vn and (norm_vn in normalize_vietnamese_name(iss.get('assignee', '')) or normalize_vietnamese_name(iss.get('assignee', '')) in norm_vn)]
        
        l_name = s_info.get("leader", "")
        if "Hùng" in l_name:
            campus_leader = "Hồ Xuân Hùng (Leader Khối CNTT Hà Nội)"
        elif "Oanh" in l_name:
            campus_leader = "Hoàng Thị Kim Oanh (Leader Khối QTKD)"
        elif "Hằng" in l_name:
            campus_leader = "Giáp Thị Minh Hằng (Leader Khối Ngoại ngữ & KNM)"
        elif "Tươi" in l_name:
            campus_leader = "Nguyễn Thị Tươi (Leader Khối QLCLĐT Hà Nội)"
        elif "Đạo" in l_name:
            campus_leader = "Nguyễn Bá Minh Đạo (Leader Khối CNTT HCM)"
        elif "Cường" in l_name:
            campus_leader = "Trần Minh Cường (Leader LMS AI)"
        else:
            campus_leader = l_name

        for d in all_dates_to_scan:
            reports_on_day = daily_cache.get(d, [])
            user_rep = None
            for r in reports_on_day:
                rep_user = normalize_name(r.get("user", ""))
                if norm_name in rep_user or rep_user in norm_name:
                    user_rep = r
                    break
                    
            if user_rep:
                # Nếu là ngày làm việc tiêu chuẩn thì tính vào reported_days_count
                if d in staff_effective_days:
                    reported_days_count += 1
                elif period_name == "sept_01_08" and reported_days_count < expected_days_count:
                    # Ghi nhận làm việc sớm vào ngày nghỉ (ví dụ 02/09)
                    reported_days_count += 1
                    
                tasks = user_rep.get("tasks", [])
                for t in tasks:
                    t_title = t.get("title", "").strip()
                    t_hours = float(t.get("hours", 0.0))
                    total_declared_hours += t_hours
                    
                    std_hours, matched_cat, is_wildcard = match_kpi_standard(
                        group, role, rank, t_title, qtkd_master, cntt_master, nn_standards, nn_mappings
                    )
                    total_standard_hours += std_hours
                    
                    dev_hours = round(t_hours - std_hours, 2)
                    dev_pct = round((dev_hours / std_hours) * 100) if std_hours > 0 else 0
                    
                    task_record = {
                        "date": d,
                        "title": t_title,
                        "declared_hours": t_hours,
                        "standard_hours": std_hours,
                        "deviation_hours": dev_hours,
                        "deviation_pct": dev_pct,
                        "matched_category": matched_cat,
                        "is_wildcard": is_wildcard,
                        "done": t.get("done", True)
                    }
                    all_reported_tasks.append(task_record)
                    
                    if not is_wildcard and t_hours >= std_hours * 1.5 and dev_hours >= 0.5:
                        excess_tasks.append({
                            **task_record,
                            "reason": f"Khai {t_hours}h so với định mức Rank {rank} là {std_hours}h (+{dev_pct}%)"
                        })
                    elif is_wildcard:
                        reason = f"Đầu việc ngoài KPI Master (Khai {t_hours}h) - Cần nghiệm thu sản phẩm"
                        if any(k in t_title.lower() for k in ["slide", "mindmap", "quiz", "video", "bài tập", "học liệu"]):
                            reason = f"Sản xuất học liệu ngoài Master (Khai {t_hours}h - Cần nghiệm thu sản phẩm)"
                        elif any(k in t_title.lower() for k in ["nghiên cứu", "rnd", "tìm hiểu", "bồi dưỡng", "ai agent"]):
                            reason = f"Nghiên cứu công nghệ/tự học (Khai {t_hours}h - Cần báo cáo kết quả)"
                        free_floating_tasks.append({
                            **task_record,
                            "reason": reason
                        })
                            
                    t_title_norm = t_title.lower()
                    for iss in user_issues:
                        iss_title = iss.get('title', '').lower()
                        if t_title_norm in iss_title or iss_title in t_title_norm:
                            iss_state = str(iss.get('state', '')).upper()
                            if iss_state not in ['DONE', 'COMPLETED', 'HOÀN THÀNH', 'CANCELLED']:
                                unverified_tasks.append({
                                    **task_record,
                                    "issue_title": iss.get('title'),
                                    "issue_state": iss.get('state'),
                                    "reason": f"Daily Log báo hoàn thành nhưng Ticket Worklane đang ở trạng thái {iss_state}"
                                })
                            break

        reported_days_count = min(reported_days_count, expected_days_count)
        compliance_rate = round((reported_days_count / expected_days_count * 100), 1) if expected_days_count > 0 else 0.0
        excess_hours = round(max(0.0, total_declared_hours - total_standard_hours), 1)
        
        core_excess = [t for t in excess_tasks if any(k in t['title'].lower() for k in ["giảng dạy", "chuẩn bị", "soạn", "chấm", "lên lớp", "buổi", "review", "duyệt"])]
        
        if reported_days_count == 0:
            rec_cat = "CHƯA NỘP BÁO CÁO"
            rec_label = "Bỏ trống báo cáo (0%)"
            badge_class = "badge-danger"
            action_note = f"Thiếu trọn vẹn {expected_days_count}/{expected_days_count} ngày báo cáo trên Worklane. Đề xuất Leader trực tiếp nhắc nhở và trừ điểm kỷ luật tác nghiệp."
        elif len(core_excess) >= (2 if period_name == "sept_01_08" else 3) and rank >= 2:
            rec_cat = "CẦN RÀ SOÁT RANK"
            rec_label = f"Cần rà soát hạ Rank ({len(core_excess)} task nghiệp vụ vượt chuẩn)"
            badge_class = "badge-warning"
            action_note = f"Đang hưởng chế độ Rank {rank} nhưng có {len(core_excess)} task nghiệp vụ cốt lõi làm chậm hơn từ 1.5x - 3x định mức chuẩn. Đề xuất Leader phỏng vấn chuyên môn, nếu không cải thiện thì rà soát Rank."
        elif len(free_floating_tasks) >= free_tasks_threshold or excess_hours >= excess_hours_threshold:
            rec_cat = "CẢNH BÁO BÙ GIỜ"
            rec_label = f"Khai bù giờ (Dôi dư +{excess_hours}h nghi vấn)"
            badge_class = "badge-danger"
            action_note = f"Có {len(free_floating_tasks)} task tự do và tổng cộng dôi dư {excess_hours}h so với định mức KPI Master. Đề xuất chỉ nghiệm thu theo barem chuẩn, cắt giảm giờ công ảo."
        else:
            rec_cat = "CHUẨN MỰC"
            rec_label = "Khai báo chuẩn mực & Đúng hạn"
            badge_class = "badge-success"
            action_note = "Khai báo bám sát định mức KPI Master, giờ công thực chất và tuân thủ kỷ luật báo cáo tốt."
            
        results[norm_name] = {
            "name": name,
            "group": group,
            "campus": campus,
            "campus_leader": campus_leader,
            "role": role,
            "rank": rank,
            "expected_days": expected_days_count,
            "reported_days": reported_days_count,
            "compliance_rate": compliance_rate,
            "total_declared_hours": round(total_declared_hours, 1),
            "total_standard_hours": round(total_standard_hours, 1),
            "excess_hours": excess_hours,
            "excess_tasks_count": len(excess_tasks),
            "free_floating_tasks_count": len(free_floating_tasks),
            "unverified_tasks_count": len(unverified_tasks),
            "excess_tasks": excess_tasks,
            "free_floating_tasks": free_floating_tasks,
            "unverified_tasks": unverified_tasks,
            "recent_tasks": all_reported_tasks,
            "recommendation_category": rec_cat,
            "recommendation_label": rec_label,
            "badge_class": badge_class,
            "action_note": action_note
        }
        
    total_staff = len(results)
    total_exp_days = sum(s["expected_days"] for s in results.values())
    total_rep_days = sum(s["reported_days"] for s in results.values())
    total_declared_all = sum(s["total_declared_hours"] for s in results.values())
    total_std_all = sum(s["total_standard_hours"] for s in results.values())
    total_excess_all = sum(s["excess_hours"] for s in results.values())
    
    count_review_rank = len([s for s in results.values() if s["recommendation_category"] == "CẦN RÀ SOÁT RANK"])
    count_deduct_hours = len([s for s in results.values() if s["recommendation_category"] == "CẢNH BÁO BÙ GIỜ"])
    count_missing = len([s for s in results.values() if s["recommendation_category"] == "CHƯA NỘP BÁO CÁO"])
    count_compliant = len([s for s in results.values() if s["recommendation_category"] == "CHUẨN MỰC"])
    
    output_payload = {
        "period_key": period_name,
        "audit_month": period_title,
        "audit_date": audit_date_str,
        "director": "Thầy Nguyễn Duy Quang",
        "summary": {
            "total_staff": total_staff,
            "overall_compliance_rate": round(total_rep_days / total_exp_days * 100, 1) if total_exp_days > 0 else 0,
            "total_declared_hours": round(total_declared_all, 1),
            "total_standard_hours": round(total_std_all, 1),
            "total_excess_hours": round(total_excess_all, 1),
            "count_review_rank": count_review_rank,
            "count_deduct_hours": count_deduct_hours,
            "count_missing": count_missing,
            "count_compliant": count_compliant
        },
        "staff_details": results
    }
    return output_payload

def audit_all_staff_v2():
    print("=== BẮT ĐẦU KIỂM TOÁN HIỆU SUẤT WORKLANE HAI KỲ ===")
    
    # 1. Kiểm toán Kỳ Mới: 01/09 - 08/09/2026
    payload_sept = audit_period_data("sept_01_08")
    sept_file = r"data/processed/worklane_audit_sept_01_08.json"
    with open(sept_file, "w", encoding="utf-8") as f:
        json.dump(payload_sept, f, indent=2, ensure_ascii=False)
    print(f"✓ Đã xuất kiểm toán Kỳ 01/09 - 08/09/2026 vào: {sept_file}")

    # 2. Kiểm toán Kỳ Lịch sử: Tháng 08/2026
    payload_aug = audit_period_data("august")
    aug_file = r"data/processed/worklane_audit_august.json"
    with open(aug_file, "w", encoding="utf-8") as f:
        json.dump(payload_aug, f, indent=2, ensure_ascii=False)
    print(f"✓ Đã xuất kiểm toán Tháng 08/2026 vào: {aug_file}")

    # 3. Xuất file tổng hợp tương thích (Combined Detailed)
    # Mặc định summary và staff_details là kỳ mới nhất (sept_01_08)
    combined_payload = {
        **payload_sept,
        "periods": {
            "sept_01_08": payload_sept,
            "august": payload_aug
        }
    }
    
    out_file = r"data/processed/worklane_audit_detailed.json"
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(combined_payload, f, indent=2, ensure_ascii=False)
    print(f"✓ Đã cập nhật worklane_audit_detailed.json đa kỳ.")

    print("\n--- TỔNG KẾT KỲ MỚI 01/09 - 08/09/2026 ---")
    s = payload_sept['summary']
    print(f"- Tổng nhân sự: {s['total_staff']}")
    print(f"- Tỷ lệ nộp báo cáo: {s['overall_compliance_rate']}%")
    print(f"- Tổng giờ khai báo: {s['total_declared_hours']}h | Chuẩn Master: {s['total_standard_hours']}h | Dôi dư: +{s['total_excess_hours']}h")
    print(f"- Cần rà soát Rank: {s['count_review_rank']} nhân sự")
    print(f"- Cảnh báo bù giờ: {s['count_deduct_hours']} nhân sự")
    print(f"- Bỏ trống báo cáo: {s['count_missing']} nhân sự")
    print(f"- Chuẩn mực: {s['count_compliant']} nhân sự")

if __name__ == "__main__":
    audit_all_staff_v2()
