import os
import sys
import openpyxl
import shutil

sys.stdout.reconfigure(encoding='utf-8')

def sanitize_excel(file_path):
    print(f"DataSanitizer: Bắt đầu làm sạch dữ liệu tại {file_path}")
    
    if not os.path.exists(file_path):
        print(f"Lỗi: Không tìm thấy file {file_path}")
        return False
        
    # Tạo bản backup trước khi xử lý
    backup_path = file_path.replace(".xlsx", "_raw.xlsx")
    if not os.path.exists(backup_path):
        shutil.copy2(file_path, backup_path)
        print(f"DataSanitizer: Đã backup file gốc sang {backup_path}")
    
    try:
        wb = openpyxl.load_workbook(file_path)
        
        # Lặp qua tất cả các sheet
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            
            # Chuẩn hóa headers và data
            for row in range(1, sheet.max_row + 1):
                for col in range(1, sheet.max_column + 1):
                    cell = sheet.cell(row=row, column=col)
                    
                    if isinstance(cell.value, str):
                        # Cắt khoảng trắng thừa ở chuỗi, đặc biệt là tên GV
                        cell.value = cell.value.strip()
                    elif cell.value is None:
                        # Với các ô dữ liệu vi phạm (từ cột 4 trở đi) nếu None thì gán bằng 0
                        # Tuy nhiên không gán vô tội vạ, chỉ gán ở các hàng có data
                        # Ở đây tạm thời để nguyên None nếu là header, nhưng nếu là row data thì cần suy xét
                        # Tốt nhất là các script phân tích sẽ tự handle None to 0, nhưng ta có thể hỗ trợ chuẩn hóa
                        pass
                        
        wb.save(file_path)
        wb.close()
        print(f"DataSanitizer: Làm sạch thành công file {file_path}")
        return True
    except Exception as e:
        print(f"DataSanitizer Lỗi: {str(e)}")
        return False

def main():
    print("=========================================")
    print("KHỞI CHẠY DATA SANITIZER (HARNESS LAYER)")
    print("=========================================")
    
    target_file = "data/inputs/PTIT_Chiso.xlsx"
    success = sanitize_excel(target_file)
    
    if success:
        print("DataSanitizer: Đã sẵn sàng môi trường dữ liệu sạch!")
        sys.exit(0)
    else:
        print("DataSanitizer: Gặp sự cố làm sạch dữ liệu. Các Agent phía sau có thể gặp lỗi rác dữ liệu.")
        sys.exit(1)

if __name__ == "__main__":
    main()
