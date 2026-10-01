import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import shutil
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

ORIGINAL_BACKUP = r"C:\Users\DELL\Desktop\Backup\PTIT\PTIT_Chiso.xlsx"
KS24_KS25_FILE = r"C:\Users\DELL\Desktop\Backup\PTIT\PTIT_Chiso_KS24_KS25.xlsx"
DATA_INPUT = r"data/inputs/PTIT_Chiso.xlsx"

# 1. TẠO FILE PTIT_Chiso_KS24_KS25.xlsx CHO NGƯỜI DÙNG THEO DÕI VÀ ĐIỀN CHỈ SỐ KS24 & KS25
print("1. Tạo file PTIT_Chiso_KS24_KS25.xlsx...")
wb_src = openpyxl.load_workbook(ORIGINAL_BACKUP)

# Tạo workbook mới
wb_ks24_25 = openpyxl.Workbook()
wb_ks24_25.remove(wb_ks24_25.active) # remove default sheet

ks24_ks25_sheets = [
    'KS24-JavaAdvance', 'KS24_JavaWeb', 'KS24_JWS', 'KS24_AI', 'KS24_AI_Intergration', 'KS24_AI_Microservice',
    'KS25_Javascript', 'KS25_Database', 'KS25_Python', 'KS25_Python_Web', 'KS25_Phantichthietkehethong',
    'KS25_QTKD_M103', 'KS25_QTKD_M104', 'KS25_QTKD_DTB201', 'KS25_QTKD_DTB202', 'KS25_QTKD_PRJ302',
    'KS25_QTKD_BA201', 'KS25_QTKD_MAN107', 'SKL_KS24'
]

for sname in ks24_ks25_sheets:
    if sname in wb_src.sheetnames:
        ws_from = wb_src[sname]
        ws_to = wb_ks24_25.create_sheet(title=sname)
        ws_to.views.sheetView[0].showGridLines = True
        
        # Copy rows & cells with formatting
        for r in range(1, ws_from.max_row + 1):
            for c in range(1, ws_from.max_column + 1):
                cell_from = ws_from.cell(row=r, column=c)
                if cell_from.value is not None:
                    cell_to = ws_to.cell(row=r, column=c, value=cell_from.value)
                    if cell_from.has_style:
                        cell_to.font = Font(name=cell_from.font.name, size=cell_from.font.size, bold=cell_from.font.bold, color=cell_from.font.color)
                        cell_to.fill = PatternFill(fill_type=cell_from.fill.fill_type, start_color=cell_from.fill.start_color, end_color=cell_from.fill.end_color)
                        cell_to.alignment = Alignment(horizontal=cell_from.alignment.horizontal, vertical=cell_from.alignment.vertical)
                        cell_to.border = Border(left=cell_from.border.left, right=cell_from.border.right, top=cell_from.border.top, bottom=cell_from.border.bottom)

        for col_letter, dim in ws_from.column_dimensions.items():
            if dim.width:
                ws_to.column_dimensions[col_letter].width = dim.width

# Tạo thêm sheet mẫu cho môn kế tiếp của KS25 QTKD: KS25_QTKD_BI
ws_bi = wb_ks24_25.create_sheet(title="KS25_QTKD_BI")
ws_bi.views.sheetView[0].showGridLines = True
ws_bi.cell(row=1, column=1, value="Tình hình PTIT_KS25_QTKD_BI (Business Intelligence)").font = Font(name="Calibri", size=14, bold=True, color="1F497D")
ws_bi.cell(row=3, column=1, value="STT").font = Font(name="Calibri", size=11, bold=True)
ws_bi.cell(row=3, column=2, value="Lớp").font = Font(name="Calibri", size=11, bold=True)
ws_bi.cell(row=3, column=3, value="Giảng viên/Trợ giảng").font = Font(name="Calibri", size=11, bold=True)
ws_bi.cell(row=3, column=4, value="[Ngày bắt đầu]").font = Font(name="Calibri", size=11, bold=True)
ws_bi.cell(row=4, column=4, value="Chuyên cần").font = Font(name="Calibri", size=11, bold=True)
ws_bi.cell(row=4, column=5, value="Bài tập").font = Font(name="Calibri", size=11, bold=True)
ws_bi.cell(row=4, column=6, value="Elearning").font = Font(name="Calibri", size=11, bold=True)
ws_bi.cell(row=5, column=1, value=1)
ws_bi.cell(row=5, column=2, value="HN-K25-QTKD1(46)").font = Font(name="Calibri", size=11, bold=True)
ws_bi.cell(row=5, column=3, value="Đặng Quỳnh Trang")
ws_bi.cell(row=7, column=1, value=2)
ws_bi.cell(row=7, column=2, value="HN-K25-QTKD2(42)").font = Font(name="Calibri", size=11, bold=True)
ws_bi.cell(row=7, column=3, value="Đặng Quỳnh Trang")
ws_bi.column_dimensions['A'].width = 8
ws_bi.column_dimensions['B'].width = 24
ws_bi.column_dimensions['C'].width = 28
ws_bi.column_dimensions['D'].width = 14
ws_bi.column_dimensions['E'].width = 14
ws_bi.column_dimensions['F'].width = 14

wb_ks24_25.save(KS24_KS25_FILE)
print(f"✓ Đã lưu file chuyên biệt cho KS24 & KS25 tại: {KS24_KS25_FILE}")

# 2. CẬP NHẬT LẠI FILE PTIT_Chiso.xlsx VỚI CHỈ SỐ CHUẨN XÁC CỦA KS26
print("\n2. Cập nhật lại các sheet KS26 với chỉ số chuẩn xác trên LMS Frontend...")

font_title = Font(name="Calibri", size=14, bold=True, color="1F497D")
font_header = Font(name="Calibri", size=11, bold=True, color="000000")
font_bold = Font(name="Calibri", size=11, bold=True)
font_regular = Font(name="Calibri", size=11)
font_date = Font(name="Calibri", size=11, bold=True, color="1F497D")

fill_header = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
fill_date = PatternFill(start_color="B8CCE4", end_color="B8CCE4", fill_type="solid")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

SHEET_CONFIGS_CORRECTED = [
    {
        "sheet_name": "KS26_Nhapmon_CNTT",
        "title": "Tình hình PTIT_KS26_IT108 - Nhập Môn Công Nghệ Thông Tin",
        "date_str": "28/09/2026",
        "classes": [
            {"stt": 1, "code": "HN-KS26-CNTT1(43)", "gv": "Trịnh Quốc Hai", "cc": 0.0, "bt": 0.0, "el": 4.65},
            {"stt": 2, "code": "HN-KS26-CNTT2(40)", "gv": "Trịnh Quốc Hai", "cc": None, "bt": None, "el": None},
            {"stt": 3, "code": "HN-KS26-CNTT3(42)", "gv": "Lương Quốc Tuấn", "cc": 0.0, "bt": 0.0, "el": 0.0},
            {"stt": 4, "code": "HN-KS26-CNTT4(27)", "gv": "Lương Quốc Tuấn", "cc": 0.0, "bt": 0.0, "el": 22.22},
            {"stt": 5, "code": "HCM-KS26-CNTT1(46)", "gv": "Lê Hà Thanh Sang", "cc": 0.0, "bt": 0.0, "el": 6.52},
            {"stt": 6, "code": "HCM-KS26-CNTT2(45)", "gv": "Lê Hà Thanh Sang", "cc": 0.0, "bt": 0.0, "el": 8.89}
        ]
    },
    {
        "sheet_name": "KS26_KN_Lamviecnhom",
        "title": "Tình hình PTIT_KS26_SKL01 - Kỹ năng làm việc nhóm",
        "date_str": "28/09/2026",
        "classes": [
            {"stt": 1, "code": "HN-KS26-CNTT1(43)", "gv": "Hoàng Thị Hậu", "cc": 4.65, "bt": 0.0, "el": 16.28},
            {"stt": 2, "code": "HN-KS26-CNTT2(40)", "gv": "Hoàng Thị Hậu", "cc": 2.50, "bt": 0.0, "el": 5.00},
            {"stt": 3, "code": "HN-KS26-CNTT3(42)", "gv": "Hoàng Thị Hậu", "cc": None, "bt": None, "el": None},
            {"stt": 4, "code": "HCM-KS26-CNTT1(46)", "gv": "Lê Nhựt Mi", "cc": None, "bt": None, "el": None},
            {"stt": 5, "code": "HCM-KS26-CNTT2(45)", "gv": "Lê Nhựt Mi", "cc": None, "bt": None, "el": None}
        ]
    },
    {
        "sheet_name": "KS26_Basic_Speaking",
        "title": "Tình hình PTIT_KS26_ENG105 - Basic Speaking",
        "date_str": "28/09/2026",
        "classes": [
            {"stt": 1, "code": "HCM-K26-QTKD1(23)", "gv": "Huỳnh Thị Kim Khánh", "cc": 0.0, "bt": 0.0, "el": 30.43},
            {"stt": 2, "code": "HCM-KS26-CNTT2(45)", "gv": "Huỳnh Thị Kim Khánh", "cc": 0.0, "bt": 0.0, "el": 11.11},
            {"stt": 3, "code": "HN-K26-QTKD3(44)", "gv": "Nguyễn Hồng Nhung", "cc": 0.0, "bt": 0.0, "el": 0.0},
            {"stt": 4, "code": "HN-KS26-CNTT3(42)", "gv": "Lò Thị Ngọc Anh", "cc": 0.0, "bt": 0.0, "el": 0.0},
            {"stt": 5, "code": "HCM-KS26-CNTT1(46)", "gv": "Huỳnh Thị Kim Khánh", "cc": None, "bt": None, "el": None},
            {"stt": 6, "code": "HN-K26-QTKD1(44)", "gv": "Lò Thị Ngọc Anh", "cc": None, "bt": None, "el": None},
            {"stt": 7, "code": "HN-K26-QTKD2(44)", "gv": "Nguyễn Hồng Nhung", "cc": None, "bt": None, "el": None},
            {"stt": 8, "code": "HN-KS26-CNTT2(40)", "gv": "Nguyễn Hồng Nhung", "cc": None, "bt": None, "el": None}
        ]
    },
    {
        "sheet_name": "KS26_QTKD_Tuduyphantich",
        "title": "Tình hình PTIT_KS26_SSK103 - Tư duy phân tích",
        "date_str": "28/09/2026",
        "classes": [
            {"stt": 1, "code": "HN-K26-QTKD1(44)", "gv": "Nguyễn Ngọc Vân Khanh", "cc": 9.09, "bt": 0.0, "el": 6.82},
            {"stt": 2, "code": "HN-K26-QTKD2(44)", "gv": "Nguyễn Thị Hồng Minh", "cc": None, "bt": None, "el": None},
            {"stt": 3, "code": "HN-K26-QTKD3(44)", "gv": "Nguyễn Ngọc Vân Khanh", "cc": None, "bt": None, "el": None},
            {"stt": 4, "code": "HCM-K26-QTKD1(23)", "gv": "Lê Thị Bảo Yến", "cc": None, "bt": None, "el": None}
        ]
    },
    {
        "sheet_name": "KS26_QTKD_Tinhocungdung",
        "title": "Tình hình PTIT_KS26_SSK102 - Tin học ứng dụng",
        "date_str": "28/09/2026",
        "classes": [
            {"stt": 1, "code": "HCM-K26-QTKD1(23)", "gv": "Lê Hà Thanh Sang", "cc": 8.70, "bt": 0.0, "el": 8.70},
            {"stt": 2, "code": "HN-K26-QTKD3(44)", "gv": "Nguyễn Thị Hồng Minh", "cc": 18.18, "bt": 0.0, "el": 9.09},
            {"stt": 3, "code": "HN-K26-QTKD1(44)", "gv": "Nguyễn Ngọc Vân Khanh", "cc": None, "bt": None, "el": None},
            {"stt": 4, "code": "HN-K26-QTKD2(44)", "gv": "Nguyễn Ngọc Vân Khanh", "cc": None, "bt": None, "el": None}
        ]
    }
]

def apply_sheet(wb, cfg):
    sname = cfg["sheet_name"]
    if sname in wb.sheetnames:
        ws = wb[sname]
        wb.remove(ws)
    ws = wb.create_sheet(title=sname)
    ws.views.sheetView[0].showGridLines = True

    ws.cell(row=1, column=1, value=cfg["title"]).font = font_title

    ws.cell(row=3, column=1, value="STT").font = font_header
    ws.cell(row=3, column=1).alignment = align_center
    ws.cell(row=3, column=1).fill = fill_header

    ws.cell(row=3, column=2, value="Lớp").font = font_header
    ws.cell(row=3, column=2).alignment = align_left
    ws.cell(row=3, column=2).fill = fill_header

    ws.cell(row=3, column=3, value="Giảng viên/Trợ giảng").font = font_header
    ws.cell(row=3, column=3).alignment = align_left
    ws.cell(row=3, column=3).fill = fill_header

    ws.cell(row=3, column=4, value=cfg["date_str"]).font = font_date
    ws.cell(row=3, column=4).alignment = align_center
    ws.cell(row=3, column=4).fill = fill_date

    ws.cell(row=4, column=4, value="Chuyên cần").font = font_bold
    ws.cell(row=4, column=4).alignment = align_center
    ws.cell(row=4, column=4).fill = fill_header

    ws.cell(row=4, column=5, value="Bài tập").font = font_bold
    ws.cell(row=4, column=5).alignment = align_center
    ws.cell(row=4, column=5).fill = fill_header

    ws.cell(row=4, column=6, value="Elearning").font = font_bold
    ws.cell(row=4, column=6).alignment = align_center
    ws.cell(row=4, column=6).fill = fill_header

    for c in range(1, 7):
        for r in [3, 4]:
            ws.cell(row=r, column=c).border = thin_border

    cur_row = 5
    for item in cfg["classes"]:
        ws.cell(row=cur_row, column=1, value=item["stt"]).font = font_regular
        ws.cell(row=cur_row, column=1).alignment = align_center

        ws.cell(row=cur_row, column=2, value=item["code"]).font = font_bold
        ws.cell(row=cur_row, column=2).alignment = align_left

        ws.cell(row=cur_row, column=3, value=item["gv"]).font = font_regular
        ws.cell(row=cur_row, column=3).alignment = align_left

        c_val = item["cc"]
        b_val = item["bt"]
        e_val = item["el"]

        cell_c = ws.cell(row=cur_row, column=4, value=c_val if c_val is not None else "")
        cell_c.font = font_regular
        cell_c.alignment = align_right

        cell_b = ws.cell(row=cur_row, column=5, value=b_val if b_val is not None else "")
        cell_b.font = font_regular
        cell_b.alignment = align_right

        cell_e = ws.cell(row=cur_row, column=6, value=e_val if e_val is not None else "")
        cell_e.font = font_regular
        cell_e.alignment = align_right

        for c in range(1, 7):
            ws.cell(row=cur_row, column=c).border = thin_border
            ws.cell(row=cur_row + 1, column=c).border = thin_border

        cur_row += 2

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 28
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 14

for target_file in [ORIGINAL_BACKUP, DATA_INPUT]:
    if os.path.exists(target_file):
        wb = openpyxl.load_workbook(target_file)
        for cfg in SHEET_CONFIGS_CORRECTED:
            apply_sheet(wb, cfg)
        # Aliases for KS25 QTKD
        aliases = [
            ("KS25_QTKD_Tuduyphantich", SHEET_CONFIGS_CORRECTED[3]),
            ("KS25_QTKD_Tinhocungdung", SHEET_CONFIGS_CORRECTED[4])
        ]
        for a_name, orig in aliases:
            a_cfg = dict(orig)
            a_cfg["sheet_name"] = a_name
            apply_sheet(wb, a_cfg)
        wb.save(target_file)
        print(f"✓ Đã cập nhật xong số liệu chuẩn xác vào: {target_file}")

print("\n🎉 HOÀN TẤT BƯỚC 1 VÀ BƯỚC 2!")
